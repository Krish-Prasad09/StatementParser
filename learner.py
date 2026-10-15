"""
learner.py
----------
One-shot Format Learner with Schema Validation, PII Redaction, and Self-Check at Learn Time.
Called ONLY the first time a new statement format is seen.
Produces a declarative reusable template (FormatTemplate) and validates it
before saving to templates/<format_id>.json.
Subsequent statements of this format are parsed deterministically with 0 LLM calls.
"""

import os
import re
import json
import csv
import hashlib
from typing import Dict, Any, Tuple, Optional
import requests
import pdfplumber
import openpyxl

from models import FormatTemplate
from parser import StatementParser
from validator import FinancialValidator


# Generic data leak detectors (rejects templates containing baked-in client personal data)
GENERIC_DATA_LEAK_PATTERNS = [
    r"\b[6-9]\d{9}\b",                                        # 10-digit mobile numbers
    r"\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Z|a-z]{2,}\b",   # Email addresses
    r"\b[A-Z]{5}[0-9]{4}[A-Z]\b",                             # Real or synthetic PANs
    r"\bIN[EF][A-Z0-9]{8}[0-9]\b",                            # Specific 12-char ISINs (INE/INF)
    r"(?i)\bRs\.?\s*[0-9]{1,3}(?:,[0-9]{2,3})+\b",            # Specific formatted rupee amounts
]

WORKED_EXAMPLE_TEMPLATE = {
    "format_id": "example_pms_monthly",
    "version": 1,
    "layout_kind": "ruled_table",
    "provider_key": "example_pms",
    "sebi_regn": "INP000000000",
    "fingerprint": {
        "summary_headers": ["Opening Corpus", "Inflows", "Outflows", "Net Gain", "Fees", "Cash", "Closing Corpus"],
        "holdings_headers": ["Security", "ISIN", "Qty", "Avg Cost", "Market Price", "Market Value", "Weight"],
        "transactions_headers": ["Date", "Type", "Particulars", "Qty", "Price", "Amount"]
    },
    "fields": {
        "provider": "Example Wealth PMS",
        "client_id_pattern": "CLIENT ID\\s*[:\\s]\\s*([A-Za-z0-9\\/\\-]+)",
        "client_name_pattern": "Client Name\\s*[:\\s]\\s*([^\\n\\r]+)",
        "pan_pattern": "PAN\\s*[:\\s]\\s*([A-Z]{5}[0-9]{4}[A-Z])",
        "strategy_pattern": "Strategy\\s*[:\\s]\\s*([^\\n\\r]+)"
    },
    "period": {
        "pattern": "Period:\\s*([0-9A-Za-z\\-\\/]+)\\s+to\\s+([0-9A-Za-z\\-\\/]+)",
        "format": "DD-Mon-YYYY"
    },
    "summary": {
        "has_gain": True,
        "mapping": {
            "opening": "Opening Corpus",
            "inflow": "Inflows",
            "outflow": "Outflows",
            "gain": "Net Gain",
            "fee": "Fees",
            "cash": "Cash",
            "closing": "Closing Corpus"
        }
    },
    "holdings": {
        "mapping": {
            "name": "Security",
            "isin": "ISIN",
            "quantity": "Qty",
            "avg_cost": "Avg Cost",
            "price": "Market Price",
            "market_value": "Market Value",
            "weight_pct": "Weight"
        },
        "total_row_labels": ["Total Equity"],
        "cash_row_labels": ["Cash Balance", "Cash"],
        "nav_row_labels": ["Closing Corpus"],
        "supports_short": False,
        "is_unlisted": False
    },
    "transactions": {
        "mapping": {
            "date": "Date",
            "raw_type": "Type",
            "security": "Particulars",
            "quantity": "Qty",
            "price": "Price",
            "amount": "Amount"
        },
        "type_map": {
            "INFLOW": "inflow",
            "OUTFLOW": "outflow",
            "BUY": "buy",
            "SELL": "sell",
            "DIVIDEND": "dividend",
            "FEE": "fee"
        }
    }
}


def extract_and_redact_document_text(sample_file: str) -> str:
    """
    Extracts raw text from a PDF, XLSX, or CSV sample file and masks all PII
    (client names, PANs, phone numbers, email addresses, client IDs, addresses, and numbers)
    before building the LLM prompt.
    """
    ext = os.path.splitext(sample_file)[1].lower()
    raw_text = ""

    try:
        if ext == ".pdf":
            with pdfplumber.open(sample_file) as pdf:
                for page in pdf.pages[:3]:
                    raw_text += (page.extract_text() or "") + "\n"
        elif ext == ".xlsx":
            wb = openpyxl.load_workbook(sample_file, data_only=True)
            for sname in wb.sheetnames:
                ws = wb[sname]
                raw_text += f"--- SHEET: {sname} ---\n"
                for r in list(ws.iter_rows(values_only=True))[:15]:
                    row_str = " | ".join(str(c) for c in r if c is not None)
                    if row_str.strip():
                        raw_text += row_str + "\n"
        elif ext == ".csv":
            with open(sample_file, "r", encoding="utf-8") as f:
                rdr = list(csv.reader(f))
            for r in rdr[:25]:
                if any(r):
                    raw_text += " , ".join(r) + "\n"
    except Exception:
        pass

    # Redact PII from extracted text
    redacted = raw_text
    # 1. Mask PANs
    redacted = re.sub(r"\b[A-Z]{5}[0-9]{4}[A-Z]\b", "<REDACTED_PAN>", redacted)
    # 2. Mask emails
    redacted = re.sub(r"\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Z|a-z]{2,}\b", "<REDACTED_EMAIL>", redacted)
    # 3. Mask phone numbers
    redacted = re.sub(r"(?:\+91[\-\s]?)?[6-9]\d{4}[\-\s]?\d{5}\b", "<REDACTED_PHONE>", redacted)
    # 4. Mask Client IDs after labels
    cid_labels = r"(?:CLIENT\s*ID|CLIENT\s*CODE|FOLIO(?:\s*/\s*CLIENT\s*CODE)?|ACCOUNT\s*NO|ACCOUNT\s*NUMBER|INVESTOR\s*CODE|UNITHOLDER\s*ID|FOLIO\s*NO|FOLIO)"
    redacted = re.sub(rf"({cid_labels}\s*[:\-]?\s*)([A-Za-z0-9\/\-]+)", r"\1<REDACTED_CLIENT_ID>", redacted, flags=re.IGNORECASE)
    # 5. Mask standalone Client ID tokens (e.g. AUR-100201, EVG-AIF-4410, CP/PMS/0017, DR-55012, HPS-3301)
    redacted = re.sub(r"\b(?!(?:INP|IN/AIF|INZ|INE|INF|IN9))[A-Z]{2,4}[-/](?:PMS/|UH/|AIF/)?[0-9]{4,6}[A-Za-z0-9\-]*\b", "<REDACTED_CLIENT_ID>", redacted)
    # 6. Mask client names after labels (including bare 'Name')
    name_labels = r"(?:Client\s*Name|Investor\s*Name|Account\s*Name|Account\s*Holder|Investor(?!\s+Code)|Name|Unitholder(?!\s+ID))"
    redacted = re.sub(rf"({name_labels}\s*[:\s|,-]\s*)([A-Z][a-zA-Z]+(?:\s+[A-Z][a-zA-Z]+)+)", r"\1<REDACTED_NAME>", redacted, flags=re.IGNORECASE)
    # 7. Mask name before PAN or delimiter
    redacted = re.sub(r"\b([A-Z][a-z]+(?:\s+[A-Z][a-z]+)+)(?=\s*[,|\t]\s*(?:<REDACTED_PAN>|[A-Z]{5}\d{4}[A-Z]|PAN))", "<REDACTED_NAME>", redacted)
    # 8. Mask addresses
    redacted = re.sub(r"(?:Address|Residing\s+at)\s*[:\s][^\n\r]+", "<REDACTED_ADDRESS>", redacted, flags=re.IGNORECASE)
    # 9. Mask monetary amounts and 5+ digit numbers
    redacted = re.sub(r"\b[0-9]{1,3}(?:,[0-9]{2,3})+(?:\.[0-9]{1,4})?\b", "<AMOUNT>", redacted)
    redacted = re.sub(r"\b\d{5,}(?:\.\d+)?\b", "<NUMBER>", redacted)
    return redacted


def _scrub_log_message(msg: str) -> str:
    if not msg:
        return ""
    s = re.sub(r"(?:key=|Bearer\s+|x-goog-api-key['\":\s=]+)[A-Za-z0-9_\.\-]{16,}", "[REDACTED_KEY]", msg)
    s = re.sub(r"\bAIza[0-9A-Za-z-_]{35}\b", "[REDACTED_KEY]", s)
    s = re.sub(r"\bAQ\.[A-Za-z0-9_\-]{40,}\b", "[REDACTED_KEY]", s)
    s = re.sub(r"\bsk-[A-Za-z0-9]{32,}\b", "[REDACTED_KEY]", s)
    for env_k in ("GEMINI_API_KEY", "OPENAI_API_KEY"):
        val = os.environ.get(env_k)
        if val and len(val) >= 8:
            s = s.replace(val, "[REDACTED_KEY]")
    return s


def extract_first_json_object(raw_text: str) -> str:
    text = raw_text.strip()
    m = re.search(r"```(?:json)?\s*([\s\S]*?)\s*```", text)
    if m:
        text = m.group(1).strip()
    start = text.find("{")
    if start == -1:
        raise ValueError("No JSON object '{' found in response")
    depth = 0
    in_string = False
    escape = False
    for i in range(start, len(text)):
        ch = text[i]
        if escape:
            escape = False
            continue
        if ch == "\\":
            escape = True
            continue
        if ch == '"':
            in_string = not in_string
            continue
        if not in_string:
            if ch == "{":
                depth += 1
            elif ch == "}":
                depth -= 1
                if depth == 0:
                    return text[start:i+1]
    return text[start:]


def append_learning_log(format_id: str, outcome: str):
    try:
        os.makedirs("logs", exist_ok=True)
        from datetime import datetime, timezone
        now_iso = datetime.now(timezone.utc).isoformat()
        entry = f"{now_iso} | format_id={format_id} | outcome={outcome}\n"
        with open(os.path.join("logs", "llm_learning.log"), "a", encoding="utf-8") as f:
            f.write(entry)
    except Exception:
        pass


class FormatLearner:
    """One-shot learner creating, schema-checking, and self-checking declarative templates."""

    def __init__(self, templates_dir: str = "templates", db_path: Optional[str] = "statements.db"):
        self.templates_dir = templates_dir
        self.db_path = db_path
        os.makedirs(self.templates_dir, exist_ok=True)
        self.parser = StatementParser()
        self.validator = FinancialValidator()
        self.llm_calls_made = 0
        self.actual_http_calls = 0
        self.failed_formats = set()
        self.last_failure_issue: str = "learning_failed_llm_unavailable"
        self.last_diagnostic_message: str = ""
        self._load_persisted_failures()

    def _load_persisted_failures(self):
        if not self.db_path or not os.path.exists(self.db_path):
            return
        try:
            from persistence import StatementStore
            store = StatementStore(self.db_path)
            for rec in store.get_failed_learnings():
                fid = rec.get("format_id")
                fhash = rec.get("fingerprint_hash") or ""
                if fid and fhash:
                    self.failed_formats.add((fid, fhash))
        except Exception:
            pass

    def _persist_failure(self, format_id: str, provider_key: str = "", layout_kind: str = "", fingerprint_hash: str = "", failure_issue: str = "", diagnostic_message: str = ""):
        if fingerprint_hash:
            self.failed_formats.add((format_id, fingerprint_hash))
        if self.db_path:
            try:
                from persistence import StatementStore
                store = StatementStore(self.db_path)
                store.save_failed_learning(
                    format_id=format_id,
                    provider_key=provider_key,
                    layout_kind=layout_kind,
                    fingerprint_hash=fingerprint_hash,
                    failure_issue=failure_issue,
                    diagnostic_message=diagnostic_message
                )
            except Exception:
                pass

    @staticmethod
    def extract_and_redact_document_text(file_path: str) -> str:
        return extract_and_redact_document_text(file_path)

    def has_template(self, format_id: str) -> bool:
        """Checks if a validated template already exists on disk."""
        safe_id = re.sub(r"[^a-z0-9_]", "_", str(format_id).lower()).strip("_")
        path = os.path.join(self.templates_dir, f"{safe_id}.json")
        return os.path.exists(path)

    def load_template(self, format_id: str) -> Optional[Dict[str, Any]]:
        """Loads an existing template."""
        safe_id = re.sub(r"[^a-z0-9_]", "_", str(format_id).lower()).strip("_")
        path = os.path.join(self.templates_dir, f"{safe_id}.json")
        if os.path.exists(path):
            with open(path, "r", encoding="utf-8") as f:
                return json.load(f)
        return None

    def build_learning_prompt(self, safe_id: str, sample_file: str, features: Dict[str, Any]) -> str:
        """
        Builds a rich LLM prompt containing:
        1. Redacted document text from the sample
        2. Extracted structural features
        3. The concrete FormatTemplate JSON Schema
        4. One worked example template
        5. Explicit JSON-only instruction
        """
        redacted_text = extract_and_redact_document_text(sample_file)
        schema_json = json.dumps(FormatTemplate.model_json_schema(), indent=2)
        example_json = json.dumps(WORKED_EXAMPLE_TEMPLATE, indent=2)
        clean_features = {k: v for k, v in features.items() if k != "file_path"}

        return f"""You are a financial statement extraction engineer.
Analyze the unseen PMS/AIF statement layout below and generate a declarative FormatTemplate JSON object for format_id '{safe_id}'.

1. REDACTED DOCUMENT TEXT:
{redacted_text}

2. EXTRACTED STRUCTURAL FEATURES:
{json.dumps(clean_features, indent=2)}

3. TARGET FORMATTEMPLATE JSON SCHEMA:
{schema_json}

4. WORKED EXAMPLE TEMPLATE:
{example_json}

STRICT RULES:
- Return ONLY a valid JSON object matching the FormatTemplate schema.
- NEVER include any client names, PANs, mobile numbers, ISINs, or monetary amounts in the template.
- Map the document's exact summary headers, holdings headers/regex, and transaction headers/type_map.
- In summary: if the document's summary table has no Net Gain/(Loss) column, set "has_gain": false.
- In holdings: ensure total_row_labels includes all total and subtotal labels appearing at the bottom of the table (e.g. ["Total Equity", "Subtotal Listed Equities", "Total Listed Equities"]).
- In holdings: ensure cash_row_labels includes cash labels (e.g. ["Liquid Cash Reserve", "Cash & Bank Balance"]).
- In holdings: if there is no Market Price column (only avg cost and market value), set "derive_price": true.
"""

    def validate_template_schema_and_pii(self, candidate_tmpl: Any) -> Tuple[bool, str]:
        """
        Validates the LLM output against the FormatTemplate Pydantic schema and PII leak rules.
        """
        if not isinstance(candidate_tmpl, dict) or not candidate_tmpl:
            return False, "Candidate template is empty or not a JSON object"

        # 1. PII Leak Check
        tmpl_str = json.dumps(candidate_tmpl)
        for pat in GENERIC_DATA_LEAK_PATTERNS:
            if re.search(pat, tmpl_str):
                return False, f"Template contains baked-in personal/monetary data matching '{pat}'"

        # 2. Pydantic Schema Validation
        try:
            validated_model = FormatTemplate.model_validate(candidate_tmpl)
        except Exception as e:
            return False, f"Template failed FormatTemplate schema validation: {str(e)}"

        # 3. Require non-empty extraction rules for holdings or summary
        has_holdings_rules = bool(validated_model.holdings.mapping or validated_model.holdings.regex or validated_model.holdings.sections or validated_model.holdings.sheet)
        has_summary_rules = bool(validated_model.summary.mapping or validated_model.summary.section_header or validated_model.summary.sections or validated_model.summary.sheet)
        if not (has_holdings_rules and has_summary_rules):
            return False, "Template is missing required summary or holdings extraction rules"

        return True, "Schema and PII validation passed"

    def self_check_template(self, candidate_tmpl: Dict[str, Any], sample_file: str) -> Tuple[bool, str]:
        """
        Runs rigorous self-check on the candidate template:
        1. Validates schema and ensures zero PII is baked in.
        2. Executes candidate template on the sample file via StatementParser.
        3. Audits extracted records with FinancialValidator and requires zero critical issues.
        """
        schema_ok, schema_msg = self.validate_template_schema_and_pii(candidate_tmpl)
        if not schema_ok:
            self.last_failure_issue = "learning_failed_invalid_template"
            return False, schema_msg

        try:
            records = self.parser.parse(sample_file, candidate_tmpl)
            if not records:
                self.last_failure_issue = "learning_failed_self_check"
                return False, "Template rejected: parser returned zero statement records on sample"
        except Exception as e:
            self.last_failure_issue = "learning_failed_self_check"
            return False, f"Template rejected: parser crashed during self-check: {str(e)}"

        for rec in records:
            if not rec.statement.client_id:
                self.last_failure_issue = "learning_failed_self_check"
                return False, "Template rejected: self-check failed to extract non-empty client_id"
            if not rec.statement.period_end:
                self.last_failure_issue = "learning_failed_self_check"
                return False, "Template rejected: self-check failed to extract non-empty period_end"
            if not rec.holdings and rec.summary.closing > 0 and rec.summary.cash == 0:
                self.last_failure_issue = "learning_failed_self_check"
                return False, "Template rejected: self-check extracted zero holdings for non-zero corpus"

        status, issues = self.validator.validate(sample_file, records, template=candidate_tmpl)
        blocking_issues = [i for i in issues if i not in self.validator.SOFT_ISSUES]
        if blocking_issues or status != "validated":
            self.last_failure_issue = "learning_failed_self_check"
            return False, f"Template rejected: validation issues on sample: {blocking_issues}"

        return True, "Self-check passed successfully"

    def _call_llm_api(self, prompt: str) -> Optional[Dict[str, Any]]:
        import logging
        import time

        env_file = os.path.join(os.path.dirname(os.path.abspath(__file__)), ".env")
        if os.path.exists(env_file):
            try:
                with open(env_file, "r", encoding="utf-8") as ef:
                    for eline in ef:
                        eline = eline.strip()
                        if eline and not eline.startswith("#") and "=" in eline:
                            ek, ev = eline.split("=", 1)
                            ek, ev = ek.strip(), ev.strip().strip("\"'")
                            if ek and not os.environ.get(ek):
                                os.environ[ek] = ev
            except Exception:
                pass

        gemini_key = os.environ.get("GEMINI_API_KEY")
        openai_key = os.environ.get("OPENAI_API_KEY")
        if gemini_key:
            model_name = os.environ.get("LLM_MODEL", "gemini-2.5-flash")
            headers = {
                "x-goog-api-key": gemini_key,
                "Content-Type": "application/json"
            }
            payload = {
                "contents": [{"parts": [{"text": prompt}]}],
                "generationConfig": {"response_mime_type": "application/json"}
            }
            url = f"https://generativelanguage.googleapis.com/v1beta/models/{model_name}:generateContent"
            self.actual_http_calls += 1
            try:
                resp = requests.post(url, headers=headers, json=payload, timeout=(10, 90))
                if resp.status_code == 200:
                    data = resp.json()
                    raw_text = data["candidates"][0]["content"]["parts"][0]["text"]
                    clean_json = extract_first_json_object(raw_text)
                    return json.loads(clean_json)
                else:
                    self.last_diagnostic_message = f"HTTP {resp.status_code}"
                    logging.warning(_scrub_log_message(f"Gemini API ({model_name}) request failed with status {resp.status_code}"))
            except Exception as e:
                self.last_diagnostic_message = f"transport error: {type(e).__name__}"
                logging.warning(_scrub_log_message(f"Gemini API ({model_name}) error: {type(e).__name__}"))
            return None

        if openai_key:
            model = os.environ.get("LLM_MODEL", "gpt-4o-mini")
            url = "https://api.openai.com/v1/chat/completions"
            headers = {
                "Authorization": f"Bearer {openai_key}",
                "Content-Type": "application/json"
            }
            payload = {
                "model": model,
                "messages": [
                    {"role": "system", "content": "You are a financial statement extraction template generator. Return only valid JSON adhering to FormatTemplate schema."},
                    {"role": "user", "content": prompt}
                ],
                "response_format": {"type": "json_object"}
            }
            self.actual_http_calls += 1
            try:
                resp = requests.post(url, headers=headers, json=payload, timeout=(10, 90))
                if resp.status_code == 200:
                    data = resp.json()
                    raw_text = data["choices"][0]["message"]["content"]
                    clean_json = extract_first_json_object(raw_text)
                    return json.loads(clean_json)
                else:
                    self.last_diagnostic_message = f"HTTP {resp.status_code}"
                    logging.warning(_scrub_log_message(f"OpenAI API request failed with status {resp.status_code}"))
            except Exception as e:
                self.last_diagnostic_message = f"transport error: {type(e).__name__}"
            return None

        if not gemini_key and not openai_key:
            self.last_diagnostic_message = "no API key configured"
        return None

    def learn_format(
        self,
        format_id: str,
        sample_file: str,
        features: Dict[str, Any],
        approved: bool = True,
    ) -> Tuple[Optional[Dict[str, Any]], int]:
        # Report the number of outbound requests for this format encounter only.
        self.actual_http_calls = 0
        self.llm_calls_made = 0
        prov_slug = features.get("provider_key") or "unseen_provider"
        layout_slug = features.get("layout_kind") or "ruled_table"
        if not format_id or format_id == "unknown_format":
            format_id = f"{prov_slug}_{layout_slug}"
        safe_id = re.sub(r"[^a-z0-9_]", "_", str(format_id).lower()).strip("_")

        if self.has_template(safe_id):
            tmpl = self.load_template(safe_id)
            return tmpl, 0

        fprint_hash = hashlib.sha256(json.dumps(features.get("table_headers", []), sort_keys=True).encode()).hexdigest()[:16]
        # A failed sample must not blacklist future layouts from the provider.
        # Suppression applies only to the exact structural fingerprint.
        if (safe_id, fprint_hash) in self.failed_formats:
            if self.db_path:
                try:
                    from persistence import StatementStore
                    store = StatementStore(self.db_path)
                    rec = store.has_failed_learning(safe_id, fingerprint_hash=fprint_hash)
                    if rec:
                        self.last_failure_issue = rec.get("failure_issue") or self.last_failure_issue
                        self.last_diagnostic_message = rec.get("diagnostic_message") or self.last_diagnostic_message
                except Exception:
                    pass
            return None, 0

        prompt = self.build_learning_prompt(safe_id, sample_file, features)
        candidate_tmpl = self._call_llm_api(prompt)

        if candidate_tmpl is None:
            self.last_failure_issue = "learning_failed_llm_unavailable"
            if not self.last_diagnostic_message:
                self.last_diagnostic_message = "no API key configured"
            self._persist_failure(safe_id, prov_slug, layout_slug, fprint_hash, self.last_failure_issue, self.last_diagnostic_message)
            append_learning_log(safe_id, f"failed: {self.last_diagnostic_message}")
            calls_used = self.actual_http_calls
            self.llm_calls_made = self.actual_http_calls
            return None, calls_used

        calls_used = max(1, self.actual_http_calls)
        self.llm_calls_made = calls_used

        def _apply_system_fields(tmpl_dict: Dict[str, Any]):
            tmpl_dict["format_id"] = safe_id
            tmpl_dict["approved"] = approved
            tmpl_dict["provider_key"] = features.get("provider_key") or safe_id
            tmpl_dict["sebi_regn"] = features.get("sebi_regn")
            tmpl_dict["layout_kind"] = features.get("layout_kind") or "ruled_table"
            tmpl_dict["fingerprint"] = {
                "table_headers": list(features.get("table_headers", [])),
                "summary_headers": list(features.get("summary_headers", [])),
                "holdings_headers": list(features.get("holdings_headers", [])),
                "transactions_headers": list(features.get("transactions_headers", [])),
                "section_titles": list(features.get("section_titles", []))
            }

        _apply_system_fields(candidate_tmpl)

        schema_ok, schema_err = self.validate_template_schema_and_pii(candidate_tmpl)
        self_check_ok = False
        self_check_err = ""
        if schema_ok:
            self_check_ok, self_check_err = self.self_check_template(candidate_tmpl, sample_file)

        if not schema_ok:
            self.last_failure_issue = "learning_failed_invalid_template"
            self.last_diagnostic_message = f"schema error: {schema_err}"
            self._persist_failure(safe_id, prov_slug, layout_slug, fprint_hash, self.last_failure_issue, self.last_diagnostic_message)
            append_learning_log(safe_id, f"failed: {self.last_diagnostic_message}")
            return None, calls_used

        if not self_check_ok:
            self.last_failure_issue = "learning_failed_self_check"
            self.last_diagnostic_message = f"self-check failed: {self_check_err}"
            self._persist_failure(safe_id, prov_slug, layout_slug, fprint_hash, self.last_failure_issue, self.last_diagnostic_message)
            append_learning_log(safe_id, f"failed: {self.last_diagnostic_message}")
            return None, calls_used

        out_path = os.path.join(self.templates_dir, f"{safe_id}.json")
        try:
            with open(out_path, "w", encoding="utf-8") as f:
                json.dump(candidate_tmpl, f, indent=2)
            append_learning_log(safe_id, "success")
        except Exception as e:
            self.last_failure_issue = "learning_failed_self_check"
            self.last_diagnostic_message = f"file write error: {str(e)}"
            self._persist_failure(safe_id, prov_slug, layout_slug, self.last_failure_issue, self.last_diagnostic_message)
            append_learning_log(safe_id, f"failed: {self.last_diagnostic_message}")
            return None, calls_used

        return candidate_tmpl, calls_used


LearnerEngine = FormatLearner
