"""
fingerprinter.py
----------------
Extracts structural fingerprints from documents to recognize known formats,
detect layout drift, or flag new unseen formats for the one-shot learner.
Uses structural tokens (SEBI registrations, organization headers, table shape) and NEVER client data.
"""

import os
import re
from difflib import SequenceMatcher
from typing import Optional, Tuple, Dict, Any, List
import pdfplumber
import openpyxl

from parser import normalize_provider_title, _col_matches_label, find_provider_in_text


def compute_string_similarity(str_a: str, str_b: str) -> float:
    from utils import normalise_header
    norm_a = normalise_header(str_a)
    norm_b = normalise_header(str_b)
    if not norm_a and not norm_b:
        return 1.0
    if not norm_a or not norm_b:
        return 0.0
    if norm_a == norm_b:
        return 1.0
    return SequenceMatcher(None, norm_a, norm_b, autojunk=False).ratio()


class FormatFingerprinter:
    def __init__(self, templates_dict: Optional[Dict[str, Any]] = None):
        self.templates = templates_dict or {}

    def extract_document_features(self, file_path: str, raw_text: str = "") -> Dict[str, Any]:
        """
        Extracts structural features from a document:
        - sebi_regn (SEBI Registration Number, e.g., INP000099001, IN/AIF2/00-00/0000)
        - doc_provider (Organization name extracted from top document title)
        - org_slug (Normalized slug of organization name)
        - provider_key (Derived from organization slug or SEBI Regn)
        - layout_kind (xlsx, csv_sectioned, monospace_text, ruled_lines_horizontal_only, ruled_table)
        - table_headers (List of header titles from tables or text lines)
        - section_titles
        """
        features: Dict[str, Any] = {
            "file_path": file_path,
            "sebi_regn": None,
            "doc_provider": None,
            "org_slug": None,
            "provider_key": None,
            "layout_kind": "ruled_table",
            "table_headers": [],
            "section_titles": []
        }

        ext = os.path.splitext(file_path)[1].lower()

        # 1. Excel files (.xlsx)
        if ext == ".xlsx":
            features["layout_kind"] = "xlsx"
            try:
                wb = openpyxl.load_workbook(file_path, data_only=True)
                features["table_headers"] = list(wb.sheetnames)
                first_ws = wb.active or wb[wb.sheetnames[0]]
                prov_val = ""
                for r in first_ws.iter_rows(values_only=True):
                    if r and r[0] and str(r[0]).strip():
                        prov_val = str(r[0]).strip()
                        break
                if prov_val:
                    features["doc_provider"] = normalize_provider_title(prov_val)
                    features["org_slug"] = re.sub(r"[^a-z0-9]+", "_", prov_val.lower())[:25].strip("_")
                    features["provider_key"] = features["org_slug"]

                for sname in wb.sheetnames:
                    ws = wb[sname]
                    for r in ws.iter_rows(values_only=True):
                        row_vals = [str(c).strip() for c in r if c is not None and str(c).strip()]
                        if len(row_vals) >= 3 and any(k in " ".join(row_vals).lower() for k in ["security", "date", "client id", "isin", "qty", "scrip"]):
                            features["table_headers"].append(f"{sname}: " + " | ".join(row_vals))
                            break
            except Exception:
                pass
            return features

        # 2. Sectioned CSV files (.csv)
        if ext == ".csv":
            features["layout_kind"] = "csv_sectioned"
            try:
                import csv
                with open(file_path, "r", encoding="utf-8") as f:
                    rdr = list(csv.reader(f))
                if rdr and rdr[0] and rdr[0][0]:
                    prov_val = str(rdr[0][0]).strip()
                    features["doc_provider"] = normalize_provider_title(prov_val)
                    features["org_slug"] = re.sub(r"[^a-z0-9]+", "_", prov_val.lower())[:25].strip("_")
                    features["provider_key"] = features["org_slug"]

                for r in rdr:
                    if not r or not any(r):
                        continue
                    first_cell = r[0].strip()
                    if first_cell in ["SUMMARY", "HOLDINGS", "TRANSACTIONS"]:
                        features["section_titles"].append(first_cell)
                        features["table_headers"].append(first_cell)
                    elif len(r) >= 3 and any(k in " ".join(r).lower() for k in ["security", "date", "client id", "isin", "qty", "scrip"]):
                        features["table_headers"].append(" | ".join([c.strip() for c in r if c.strip()]))
            except Exception:
                pass
            return features

        # 3. PDF Files
        sebi_match = re.search(r"\b(INP\d{6,10}|IN/AIF[123]/[0-9A-Za-z\-/]+)\b", raw_text, re.IGNORECASE)
        if sebi_match:
            features["sebi_regn"] = sebi_match.group(1).upper()

        lines = [l.strip() for l in raw_text.split("\n") if l.strip()]
        org_title = find_provider_in_text(raw_text, default_fallback="")
        if org_title:
            features["doc_provider"] = org_title
            features["org_slug"] = re.sub(r"[^a-z0-9]+", "_", org_title.lower())[:25].strip("_")

        if features["org_slug"]:
            features["provider_key"] = features["org_slug"]
        elif features["sebi_regn"]:
            features["provider_key"] = re.sub(r"[^a-z0-9]+", "_", features["sebi_regn"].lower()).strip("_")

        # Layout detection via physical line geometry from PDF canvas
        try:
            with pdfplumber.open(file_path) as pdf:
                if pdf.pages:
                    p0 = pdf.pages[0]
                    h_lines = [l for l in p0.lines if abs(l["top"] - l["bottom"]) < 2 and abs(l["x1"] - l["x0"]) > 80]
                    v_lines = [l for l in p0.lines if abs(l["x1"] - l["x0"]) < 2 and abs(l["top"] - l["bottom"]) > 15]

                    if len(v_lines) == 0 and len(h_lines) == 0:
                        features["layout_kind"] = "monospace_text"
                    elif len(v_lines) == 0 and len(h_lines) > 0:
                        features["layout_kind"] = "ruled_lines_horizontal_only"
                    else:
                        features["layout_kind"] = "ruled_table"
        except Exception:
            pass

        # Extract table headers using pdfplumber if ruled table
        if ext == ".pdf" and features["layout_kind"] == "ruled_table":
            try:
                with pdfplumber.open(file_path) as pdf:
                    total_pages = len(pdf.pages)
                    if total_pages <= 3:
                        target_indices = set(range(total_pages))
                    else:
                        target_indices = {0}
                        found_sum = False
                        found_hold = False
                        found_txn = False

                        SUMMARY_KEYS = ["opening", "inflow", "gain", "corpus", "worth at start", "closing", "opening nav", "closing nav"]
                        HOLDINGS_KEYS = ["security", "isin", "qty", "units", "scrip", "avg cost", "market price", "rate", "worth", "holding"]
                        TXN_KEYS = ["trade date", "date", "action", "type", "particulars", "net amount", "trade", "transaction"]

                        for idx, page in enumerate(pdf.pages):
                            page_text = (page.extract_text() or "").lower()
                            if not found_sum and any(k in page_text for k in SUMMARY_KEYS):
                                target_indices.add(idx)
                                found_sum = True
                            if not found_hold and any(k in page_text for k in HOLDINGS_KEYS):
                                target_indices.add(idx)
                                found_hold = True
                            if not found_txn and any(k in page_text for k in TXN_KEYS):
                                target_indices.add(idx)
                                found_txn = True
                            if found_sum and found_hold and found_txn:
                                break

                    for idx in sorted(target_indices):
                        page = pdf.pages[idx]
                        tables = page.extract_tables()
                        for t in tables:
                            if t and t[0]:
                                clean_row = [str(col).replace("\n", " ").strip() for col in t[0] if col is not None]
                                if clean_row and len(clean_row) >= 3:
                                    hdr_str = " | ".join(clean_row)
                                    if hdr_str not in features["table_headers"]:
                                        features["table_headers"].append(hdr_str)
                                    row_lower = " ".join(clean_row).lower()
                                    if any(k in row_lower for k in ["opening", "inflow", "gain", "corpus", "worth at start", "closing", "cash"]):
                                        features["summary_headers"] = list(clean_row)
                                    elif any(k in row_lower for k in ["security", "isin", "qty", "units", "scrip", "avg cost", "market price", "rate", "worth"]):
                                        features["holdings_headers"] = list(clean_row)
                                    elif any(k in row_lower for k in ["trade date", "date", "action", "type", "particulars", "net amount", "trade"]):
                                        features["transactions_headers"] = list(clean_row)
            except Exception:
                pass

        # For Monospace textflow: extract header lines
        if features["layout_kind"] == "monospace_text" and raw_text:
            for l in lines:
                norm_l = re.sub(r"\s+", " ", l).strip()
                if any(norm_l.lower().startswith(p) for p in ["# security isin", "date type particulars", "security isin qty", "# scrip isin"]):
                    features["table_headers"].append(norm_l)

        # For Horizontal Lines ledger: extract header lines
        if features["layout_kind"] == "ruled_lines_horizontal_only" and raw_text:
            for idx, l in enumerate(lines):
                if any(hdr_token in l.lower() for hdr_token in ["holdings", "ledger entries", "transactions"]):
                    cols = []
                    for next_line in lines[idx + 1:idx + 8]:
                        if any(c in next_line.lower() for c in ["instrument", "isin", "qty", "total cost", "market value"]):
                            cols.append(next_line.strip())
                    if cols:
                        features["table_headers"].append(" ".join(cols))
                        break

        features["fingerprint"] = {
            "table_headers": list(features.get("table_headers", [])),
            "summary_headers": list(features.get("summary_headers", [])),
            "holdings_headers": list(features.get("holdings_headers", [])),
            "transactions_headers": list(features.get("transactions_headers", [])),
            "section_titles": list(features.get("section_titles", []))
        }

        return features

    def match_template(self, features: Dict[str, Any]) -> Tuple[Optional[str], Optional[Any], bool]:
        """
        Matches extracted features against known format templates.
        Strictly checks SEBI registration and provider name so a different provider is never
        silently misattributed to an existing template.
        Returns:
          (matched_format_id, template_object, is_drift)
        """
        doc_sebi = (features.get("sebi_regn") or "").strip().upper()
        doc_prov = (features.get("doc_provider") or "").strip()
        doc_org_slug = (features.get("org_slug") or "").strip().lower()
        layout_kind = features.get("layout_kind")
        doc_headers = " || ".join(features.get("table_headers", []))

        # Flatten individual column headers from table_headers for granular column drift checks
        doc_col_tokens: List[str] = []
        for h_line in features.get("table_headers", []):
            if ":" in h_line and layout_kind == "xlsx":
                h_line = h_line.split(":", 1)[1]
            for col in h_line.split("|"):
                if col.strip():
                    doc_col_tokens.append(col.strip())

        matching_candidates = []

        for format_id, template in self.templates.items():
            tmpl_layout = template.get("layout_kind")
            if tmpl_layout != layout_kind:
                continue

            tmpl_sebi = (template.get("sebi_regn") or "").strip().upper()
            tmpl_prov = (template.get("fields", {}).get("provider") or "").strip()
            tmpl_prov_key = (template.get("provider_key") or "").strip().lower()

            # 1. Explicit SEBI conflict check: different SEBI numbers cannot share a template
            if doc_sebi and tmpl_sebi and doc_sebi != tmpl_sebi:
                continue

            # 2. Positive provider identity check:
            # If doc_sebi matches tmpl_sebi identically, provider matches unambiguously
            provider_matches = False
            if doc_sebi and tmpl_sebi and doc_sebi == tmpl_sebi:
                provider_matches = True
            else:
                # If no SEBI, check provider name conflict & similarity
                if doc_prov and tmpl_prov:
                    prov_sim = compute_string_similarity(doc_prov, tmpl_prov)
                    if prov_sim < 0.75:
                        continue
                if doc_prov and tmpl_prov and compute_string_similarity(doc_prov, tmpl_prov) >= 0.80:
                    provider_matches = True
                elif doc_org_slug and tmpl_prov_key and doc_org_slug == tmpl_prov_key:
                    provider_matches = True

            if not provider_matches:
                continue

            # Build template header representation
            fp = template.get("fingerprint", {})
            parts = []
            if "table_headers" in fp and fp["table_headers"]:
                parts.extend(fp["table_headers"])
            else:
                for k in ["summary_headers", "holdings_headers", "transactions_headers"]:
                    val = fp.get(k)
                    if val:
                        if isinstance(val, list):
                            parts.append(" | ".join(str(x) for x in val))
                        else:
                            parts.append(str(val))
                if "sheet_names" in fp:
                    parts.append(" | ".join(fp["sheet_names"]))
                if "sections" in fp:
                    parts.append(" | ".join(fp["sections"]))

            tmpl_headers = " || ".join(parts)
            sim = compute_string_similarity(doc_headers, tmpl_headers)

            # Check if any mapped column in holdings (or summary for ruled_table) is missing from doc_col_tokens
            missing_mapped_col = False
            if layout_kind in ("ruled_table", "xlsx", "csv_sectioned") and doc_col_tokens:
                expected_cols = list(template.get("holdings", {}).get("mapping", {}).values())
                if layout_kind == "ruled_table":
                    expected_cols += list(template.get("summary", {}).get("mapping", {}).values())
                for exp_col in expected_cols:
                    if exp_col and not any(_col_matches_label(dc, exp_col) for dc in doc_col_tokens):
                        missing_mapped_col = True
                        break

            is_approved = template.get("approved", True)
            if template.get("status") == "pending_approval":
                is_approved = False

            # Clean match requires similarity >= 0.75, no missing mapped columns, and approved status
            clean_match = bool(sim >= 0.75 and not missing_mapped_col and is_approved)
            score = (sim + 1.0) if clean_match else sim

            matching_candidates.append({
                "format_id": format_id,
                "template": template,
                "score": score,
                "sim": sim,
                "missing_mapped_col": missing_mapped_col,
                "is_approved": is_approved,
                "clean_match": clean_match
            })

        if not matching_candidates:
            return None, None, False

        # Select the best-scoring candidate among all products/versions of this provider
        matching_candidates.sort(key=lambda c: c["score"], reverse=True)
        best = matching_candidates[0]

        if best["clean_match"]:
            return best["format_id"], best["template"], False
        else:
            return best["format_id"], best["template"], True

        return None, None, False
